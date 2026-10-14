"""Cek kelengkapan dokumen bahan baku untuk produk DR. WIDYCLINIC.

Membaca matriks bahan baku -> produk dari file Excel sumber (mis.
"HALAL - NA - DR. WIDYCLINIC (update).xlsx"), lalu mencocokkannya dengan data
di Supabase (raw_materials, raw_material_company_docs, raw_material_batches)
dan menghasilkan file .xlsx laporan kelengkapan dokumen per bahan baku.

Dokumen yang dicek (mengikuti checklist Folder ZIP Bab II aplikasi):
    1. Spesifikasi       -> raw_material_company_docs.spec_sheet_file_url
                            (PDF) ATAU parameter spek terisi teks
    2. MSDS              -> raw_material_company_docs.msds_file_url
    3. CoA               -> raw_material_batches.coa_file_url (batch terbaru)
    4. Sertifikat Halal  -> raw_material_batches.halal_batch_file_url (batch terbaru)
    5. Catatan Pemeriksaan -> raw_material_batches.qc_report_file_url (PDF)
                            ATAU hasil uji QC aktual terisi (batch terbaru)

Batch "terbaru" mengikuti logika yang sama dengan tab Cek Kelengkapan Dokumen
di halaman /raw-materials: batch dengan kesimpulan "lab" di-skip, sisanya
diurutkan created_at desc dan yang teratas yang dipakai.

Perusahaan dicek sesuai profil produk di Excel (semua produk DR. WIDYCLINIC
di database berstatus PT Erfi).

Produk yang tidak ada di file Excel sumber tapi sudah terdaftar + punya formula
di sistem bisa ditambahkan lewat EXTRA_PRODUKSI_DARI_DB di bawah; bahan bakunya
ditarik dari product_formula_lines. Bahan yang sudah terdaftar di Excel cukup
ditandai dipakai produk baru, bahan yang belum ada ditambahkan sebagai baris baru.

Pakai:
    python scripts/cek_kelengkapan_dokumen_bahan.py [file_sumber.xlsx]
"""

from __future__ import annotations

import os
import sys
from datetime import datetime
from difflib import SequenceMatcher
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WIB = ZoneInfo("Asia/Jakarta")
APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

DEFAULT_SOURCE = os.path.join(
    APP_DIR, "excel-example", "HALAL - NA - DR. WIDYCLINIC (update).xlsx"
)
DEFAULT_OUTPUT = os.path.join(
    APP_DIR, "excel-example", "CEK KELENGKAPAN DOKUMEN BAHAN BAKU - DR. WIDYCLINIC.xlsx"
)
COMPANY = "PT Erfi"

# Produk tambahan yang bahannya ditarik dari formula di sistem (bukan dari file
# Excel sumber). Nama harus cocok dengan nama produk di database.
EXTRA_PRODUKSI_DARI_DB = ["DR. WIDYCLINIC All Skin Type New Facial Wash"]

DOC_LABELS = [
    ("spec", "Spesifikasi"),
    ("msds", "MSDS"),
    ("coa", "CoA"),
    ("halal", "Sertifikat Halal"),
    ("qc", "Catatan Pemeriksaan"),
]

# ==================== STYLING (mengikuti app/excel_generator.py) ====================
# PT Erfi -> peach (#FFCC99), sama seperti header dokumen formula PT Erfi.
HEADER_FILL = PatternFill("solid", fgColor="FFCC99")
TITLE_FONT = Font(name="Arial", size=14, bold=True)
SUBTITLE_FONT = Font(name="Arial", size=11, italic=True, color="444444")
INFO_LABEL_FONT = Font(name="Arial", size=11, bold=True)
INFO_VALUE_FONT = Font(name="Arial", size=11)
HEADER_FONT = Font(name="Arial", size=11, bold=True)
BODY_FONT = Font(name="Trebuchet MS", size=11)
BODY_BOLD_FONT = Font(name="Trebuchet MS", size=11, bold=True)
NOTE_FONT = Font(name="Arial", size=10, italic=True, color="555555")

_SIDE_THIN = Side(style="thin", color="000000")
BORDER_THIN = Border(left=_SIDE_THIN, right=_SIDE_THIN, top=_SIDE_THIN, bottom=_SIDE_THIN)

# Warna status klasik Excel: hijau = ada/lengkap, merah = belum ada.
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
OK_FONT = Font(name="Trebuchet MS", size=11, bold=True, color="006100")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
BAD_FONT = Font(name="Trebuchet MS", size=11, bold=True, color="9C0006")
NA_FILL = PatternFill("solid", fgColor="E7E6E6")
NA_FONT = Font(name="Trebuchet MS", size=11, bold=True, color="595959")
SUMMARY_FILL = PatternFill("solid", fgColor="F2F2F2")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_CENTER_NOWRAP = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)


# ==================== PARSING EXCEL SUMBER ====================
def _norm(text) -> str:
    return " ".join(str(text or "").strip().lower().split())


def read_source_matrix(path: str):
    """Baca daftar bahan baku + matriks pemakaian per produk dari file sumber.

    Return: (products: dict[int, str], materials: list[dict])
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    # Baris berisi nomor produk (1..17) -> kolom produk apa di bawah header.
    number_row = None
    product_cols = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        nums = [c.value for c in row if isinstance(c.value, (int, float))]
        if len(nums) >= 5 and all(isinstance(v, (int, float)) for v in nums):
            number_row = row[0].row
            break
    if number_row is None:
        raise RuntimeError("Baris nomor produk (1, 2, 3, ...) tidak ditemukan di file sumber.")

    max_num = 0
    for col in range(3, ws.max_column + 1):
        val = ws.cell(row=number_row, column=col).value
        if isinstance(val, (int, float)):
            product_cols[int(val)] = col
            max_num = max(max_num, int(val))
    if not product_cols:
        raise RuntimeError("Tidak ada kolom produk yang terbaca di file sumber.")

    # Nama produk dari kolom Keterangan ("1. Nama Produk").
    products = {}
    for row in range(number_row + 1, ws.max_row + 1):
        for col in range(3, ws.max_column + 1):
            val = ws.cell(row=row, column=col).value
            if val is None:
                continue
            text = str(val).strip()
            head = text.split(".")[0].strip()
            if head.isdigit() and "." in text and int(head) not in products:
                products[int(head)] = text.split(".", 1)[1].strip()

    materials = []
    for row in range(number_row + 1, ws.max_row + 1):
        nama = ws.cell(row=row, column=2).value
        if not nama or not str(nama).strip():
            continue
        used = sorted(
            int(num) for num, col in product_cols.items()
            if str(ws.cell(row=row, column=col).value or "").strip().upper() == "X"
        )
        materials.append({
            "no": int(ws.cell(row=row, column=1).value or len(materials) + 1),
            "nama": str(nama).strip(),
            "used_in": used,
        })

    wb.close()
    return products, materials


# ==================== DATA DARI DATABASE ====================
def fetch_raw_material_data():
    from app.database import supabase

    rm_resp = supabase.table("raw_materials").select(
        "id, nama_dagang, kode_bahan_baku, produsen"
    ).execute()
    raw_materials = rm_resp.data or []

    docs_resp = supabase.table("raw_material_company_docs").select(
        "raw_material_id, perusahaan, spec_parameters, spec_sheet_file_url, msds_file_url"
    ).execute()
    docs_by_rm = {}
    for d in (docs_resp.data or []):
        docs_by_rm.setdefault(d["raw_material_id"], {})[d.get("perusahaan")] = d

    batch_resp = supabase.table("raw_material_batches").select(
        "id, raw_material_id, perusahaan, no_batch, created_at, kesimpulan, "
        "coa_file_url, halal_batch_file_url, qc_report_file_url, qc_results"
    ).execute()
    # Batch terbaru per (bahan baku, perusahaan): skip kesimpulan "lab",
    # urutkan created_at desc — sama dengan tab Cek Kelengkapan Dokumen.
    latest_batch = {}
    for b in (batch_resp.data or []):
        if (b.get("kesimpulan") or "") == "lab":
            continue
        key = (b["raw_material_id"], b.get("perusahaan"))
        current = latest_batch.get(key)
        if current is None or (b.get("created_at") or "") > (current.get("created_at") or ""):
            latest_batch[key] = b

    return raw_materials, docs_by_rm, latest_batch


def match_material(nama: str, raw_materials):
    """Cocokkan nama bahan dari Excel ke raw_materials: exact -> substring -> fuzzy."""
    key = _norm(nama)
    by_norm = {}
    for rm in raw_materials:
        nn = _norm(rm.get("nama_dagang"))
        by_norm.setdefault(nn, rm)
        if nn and nn not in by_norm:
            by_norm[nn] = rm
    if key in by_norm:
        return by_norm[key], "exact"
    for nn, rm in by_norm.items():
        if nn and (key in nn or nn in key):
            return rm, "substring"
    best, best_ratio = None, 0.0
    for nn, rm in by_norm.items():
        ratio = SequenceMatcher(None, key, nn).ratio()
        if ratio > best_ratio:
            best, best_ratio = rm, ratio
    if best is not None and best_ratio >= 0.6:
        return best, f"fuzzy ({best_ratio:.0%})"
    return None, None


def _list_filled(value) -> bool:
    """True kalau spec_parameters/qc_results (list of dict) ada value terisi."""
    if not isinstance(value, list):
        return bool(value)
    return any(
        isinstance(item, dict) and str(item.get("value") or "").strip()
        for item in value
    )


def compute_doc_status(rm_id, docs_by_rm, latest_batch, company):
    """Hitung status 5 dokumen Bab II untuk satu bahan baku @ satu perusahaan."""
    doc = (docs_by_rm.get(rm_id) or {}).get(company) or {}
    batch = latest_batch.get((rm_id, company))

    spec_ok = bool(doc.get("spec_sheet_file_url")) or _list_filled(doc.get("spec_parameters"))
    msds_ok = bool(doc.get("msds_file_url"))
    coa_ok = bool(batch and batch.get("coa_file_url"))
    halal_ok = bool(batch and batch.get("halal_batch_file_url"))
    qc_ok = bool(batch) and (
        bool(batch.get("qc_report_file_url")) or _list_filled(batch.get("qc_results"))
    )
    return {
        "spec": spec_ok, "msds": msds_ok, "coa": coa_ok,
        "halal": halal_ok, "qc": qc_ok, "has_batch": batch is not None,
    }


# ==================== PRODUK TAMBAHAN DARI FORMULA SISTEM ====================
def add_products_from_db(products, materials, raw_materials):
    """Tambah produk ekstra + bahan bakunya dari product_formula_lines.

    Bahan yang sudah ada di Excel (dicocokkan lewat id bahan baku sistem) cukup
    ditandai dipakai produk baru; bahan yang belum ada ditambahkan sebagai
    baris baru. Persentase formula tidak ditampilkan — laporan ini memakai
    matriks pemakaian + kelengkapan dokumen.

    Return: (jumlah_produk_tambahan, jumlah_bahan_baru).
    """
    from app.database import supabase

    added_products = 0
    added_materials = 0
    rm_by_id = {rm["id"]: rm for rm in raw_materials}
    entry_by_rm_id = {m["rm"]["id"]: m for m in materials if m.get("rm")}
    next_no = max((m["no"] for m in materials), default=0) + 1

    for name in EXTRA_PRODUKSI_DARI_DB:
        prod_resp = supabase.table("products") \
            .select("id, nama_produk, perusahaan, is_deleted") \
            .ilike("nama_produk", name).eq("is_deleted", False).limit(1).execute()
        prod = (prod_resp.data or [None])[0]
        if not prod:
            print(f"[PERINGATAN] Produk tambahan '{name}' tidak ditemukan di sistem - dilewati.")
            continue

        num = max(products) + 1 if products else 1
        products[num] = prod["nama_produk"]

        lines_resp = supabase.table("product_formula_lines") \
            .select("raw_material_id") \
            .eq("product_id", prod["id"]).execute()
        line_ids = [ln.get("raw_material_id") for ln in (lines_resp.data or []) if ln.get("raw_material_id")]

        if not line_ids:
            print(f"[PERINGATAN] Produk '{prod['nama_produk']}' belum punya formula - tanpa bahan.")
            added_products += 1
            continue

        for rid in dict.fromkeys(line_ids):  # unik, urutan formula terjaga
            if rid in entry_by_rm_id:
                entry = entry_by_rm_id[rid]
                if num not in entry["used_in"]:
                    entry["used_in"] = sorted(entry["used_in"] + [num])
                continue
            rm = rm_by_id.get(rid)
            if not rm:
                continue
            materials.append({
                "no": next_no,
                "nama": rm.get("nama_dagang") or "-",
                "used_in": [num],
                "rm": rm,
                "match_type": "exact",
                "from_db": True,
                "status": None,
                "missing": [],
                "complete": False,
            })
            entry_by_rm_id[rid] = materials[-1]
            next_no += 1
            added_materials += 1
        added_products += 1

    return added_products, added_materials


# ==================== GENERASI XLSX ====================
def _page_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = 9  # A4
    ws.sheet_view.zoomScale = 90


def _title_block(ws, last_col: int, title: str, subtitle: str):
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_col)
    cell = ws.cell(row=1, column=2, value=title)
    cell.font = TITLE_FONT
    cell.alignment = ALIGN_CENTER
    ws.row_dimensions[1].height = 22

    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    cell = ws.cell(row=2, column=2, value=subtitle)
    cell.font = SUBTITLE_FONT
    cell.alignment = ALIGN_CENTER
    ws.row_dimensions[2].height = 16


def _info_block(ws, start_row: int, pairs):
    row = start_row
    for label, value in pairs:
        label_cell = ws.cell(row=row, column=2, value=label)
        label_cell.font = INFO_LABEL_FONT
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        value_cell = ws.cell(row=row, column=3, value=value)
        value_cell.font = INFO_VALUE_FONT
        value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 15
        row += 1
    return row


def _table_header(ws, row: int, headers, first_col: int = 2):
    for offset, title in enumerate(headers):
        cell = ws.cell(row=row, column=first_col + offset, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN
    ws.row_dimensions[row].height = 30
    return row + 1


def _status_cell(ws, row: int, col: int, ok, label_ok="Ada", label_bad="Belum"):
    cell = ws.cell(row=row, column=col, value=label_ok if ok else label_bad)
    cell.alignment = ALIGN_CENTER_NOWRAP
    cell.border = BORDER_THIN
    if ok:
        cell.fill = OK_FILL
        cell.font = OK_FONT
    else:
        cell.fill = BAD_FILL
        cell.font = BAD_FONT
    return cell


def _body_cell(ws, row: int, col: int, value, align=None, bold=False, font=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font or (BODY_BOLD_FONT if bold else BODY_FONT)
    cell.alignment = align or ALIGN_LEFT
    cell.border = BORDER_THIN
    return cell


def build_report(source_path: str, output_path: str):
    products, materials = read_source_matrix(source_path)
    raw_materials, docs_by_rm, latest_batch = fetch_raw_material_data()

    # Cocokkan tiap bahan Excel ke bahan baku sistem LEBIH DULU (hasilnya juga
    # dipakai add_products_from_db supaya bahan yang sama tidak terdaftar dua kali).
    for m in materials:
        rm, match_type = match_material(m["nama"], raw_materials)
        m["rm"] = rm
        m["match_type"] = match_type

    # Produk tambahan dari formula sistem (bahan baru ikut masuk daftar).
    n_extra_prod, n_extra_mat = add_products_from_db(products, materials, raw_materials)

    # --- Hitung kelengkapan dokumen per bahan ---
    results = []
    for m in materials:
        rm = m.get("rm")
        entry = {
            "no": m["no"], "nama": m["nama"], "used_in": m["used_in"],
            "rm": rm, "match_type": m.get("match_type"), "status": None,
        }
        if rm:
            status = compute_doc_status(rm["id"], docs_by_rm, latest_batch, COMPANY)
            entry["status"] = status
            entry["missing"] = [label for key, label in DOC_LABELS if not status[key]]
            entry["complete"] = not entry["missing"]
        else:
            entry["missing"] = []
            entry["complete"] = False
        results.append(entry)

    total = len(results)
    n_complete = sum(1 for r in results if r["complete"])
    n_partial = sum(1 for r in results if r["rm"] and not r["complete"])
    n_unregistered = sum(1 for r in results if not r["rm"])
    per_doc_missing = {
        label: sum(1 for r in results if r["rm"] and not r["status"][key])
        for key, label in DOC_LABELS
    }
    now_str = datetime.now(WIB).strftime("%d-%m-%Y %H:%M WIB")

    wb = Workbook()

    # ==================== SHEET 1: CEK KELENGKAPAN DOKUMEN ====================
    ws = wb.active
    ws.title = "Cek Kelengkapan Dokumen"
    last_col = 13  # B..M
    sumber = os.path.basename(source_path)
    if n_extra_prod:
        sumber += f" + formula sistem ({n_extra_prod} produk tambahan, {n_extra_mat} bahan baru)"
    _title_block(
        ws, last_col,
        "CEK KELENGKAPAN DOKUMEN BAHAN BAKU",
        f"Produk DR. WIDYCLINIC ({len(products)} produk) - {COMPANY} - Sumber: {sumber}",
    )

    ringkas = ", ".join(f"{label} {per_doc_missing[label]}/{total}" for _, label in DOC_LABELS)
    info_pairs = [
        ("Tanggal Cek", f": {now_str}"),
        ("Perusahaan", f": {COMPANY} (sesuai profil produk di sistem)"),
        ("Jumlah Produk", f": {len(products)}"),
        ("Jumlah Bahan Baku", f": {total}"),
        ("Dokumen Lengkap (5/5)", f": {n_complete} bahan"),
        ("Dokumen Belum Lengkap", f": {n_partial} bahan"),
        ("Belum Terdaftar di Sistem", f": {n_unregistered} bahan"),
        ("Dokumen yang Belum Ada", f": {ringkas}"),
    ]
    row = _info_block(ws, 4, info_pairs)

    row += 1
    note = ws.cell(
        row=row, column=2,
        value="Catatan: file dinyatakan \"Lengkap\" bila kelima dokumen Bab II tersedia. "
              "Dokumen batch (CoA, Sertifikat Halal, Catatan Pemeriksaan) dinilai dari batch "
              "terbaru bahan baku @ perusahaan, mengikuti logika tab Cek Kelengkapan Dokumen di aplikasi.",
    )
    note.font = NOTE_FONT
    note.alignment = ALIGN_LEFT_TOP
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
    ws.row_dimensions[row].height = 28

    header_row = row + 2
    row = _table_header(ws, header_row, [
        "No", "Nama Bahan", "Bahan Baku di Sistem", "Dipakai di Produk (No.)",
        "Spesifikasi", "MSDS", "CoA", "Sertifikat Halal", "Catatan Pemeriksaan",
        "Kelengkapan", "Status", "Dokumen yang Belum Ada",
    ])

    data_start = row
    for r in results:
        rm = r["rm"]
        _body_cell(ws, row, 2, r["no"], align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws, row, 3, r["nama"], bold=True)
        if rm:
            label = f"{rm.get('kode_bahan_baku') or '-'} | {rm.get('nama_dagang') or '-'}"
            if r["match_type"] != "exact":
                label += f"  [cocok via {r['match_type']}]"
        else:
            label = "BELUM TERDAFTAR DI SISTEM"
        _body_cell(ws, row, 4, label)
        _body_cell(ws, row, 5, ", ".join(str(n) for n in r["used_in"]) or "-",
                   align=ALIGN_CENTER)

        if not rm:
            for col in range(6, 11):
                cell = ws.cell(row=row, column=col, value="-")
                cell.fill = NA_FILL
                cell.font = NA_FONT
                cell.alignment = ALIGN_CENTER_NOWRAP
                cell.border = BORDER_THIN
        else:
            status = r["status"]
            _status_cell(ws, row, 6, status["spec"])
            _status_cell(ws, row, 7, status["msds"])
            _status_cell(ws, row, 8, status["coa"])
            _status_cell(ws, row, 9, status["halal"])
            _status_cell(ws, row, 10, status["qc"])

        count = sum(r["status"][key] for key, _ in DOC_LABELS) if r["rm"] else 0
        _body_cell(ws, row, 11, f"{count}/5", align=ALIGN_CENTER_NOWRAP)

        if not rm:
            status_cell = _body_cell(ws, row, 12, "Belum Terdaftar", align=ALIGN_CENTER_NOWRAP, bold=True)
            status_cell.fill = NA_FILL
            status_cell.font = NA_FONT
            _body_cell(ws, row, 13, "Bahan baku belum terdaftar di sistem")
        elif r["complete"]:
            status_cell = _body_cell(ws, row, 12, "Lengkap", align=ALIGN_CENTER_NOWRAP, bold=True)
            status_cell.fill = OK_FILL
            status_cell.font = OK_FONT
            _body_cell(ws, row, 13, "-", align=ALIGN_CENTER_NOWRAP)
        else:
            status_cell = _body_cell(ws, row, 12, "Belum Lengkap", align=ALIGN_CENTER_NOWRAP, bold=True)
            status_cell.fill = BAD_FILL
            status_cell.font = BAD_FONT
            _body_cell(ws, row, 13, ", ".join(r["missing"]))
        ws.row_dimensions[row].height = 16
        row += 1
    data_end = row - 1

    # Baris rekap
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    recap = ws.cell(row=row, column=2, value=f"REKAP ({total} bahan baku)")
    recap.font = HEADER_FONT
    recap.fill = SUMMARY_FILL
    recap.alignment = ALIGN_CENTER_NOWRAP
    recap.border = BORDER_THIN
    for col in range(3, 6):
        ws.cell(row=row, column=col).border = BORDER_THIN
        ws.cell(row=row, column=col).fill = SUMMARY_FILL
    for key, label in DOC_LABELS:
        col = 6 + [k for k, _ in DOC_LABELS].index(key)
        ok_count = sum(1 for r in results if r["rm"] and r["status"][key])
        cell = ws.cell(row=row, column=col, value=f"{ok_count}/{total}")
        cell.font = HEADER_FONT
        cell.fill = SUMMARY_FILL
        cell.alignment = ALIGN_CENTER_NOWRAP
        cell.border = BORDER_THIN
    for col, value in ((11, f"{n_complete}/{total}"), (12, f"Lengkap: {n_complete}"),):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = HEADER_FONT
        cell.fill = SUMMARY_FILL
        cell.alignment = ALIGN_CENTER_NOWRAP
        cell.border = BORDER_THIN
    cell = ws.cell(row=row, column=13,
                   value=f"Belum lengkap: {n_partial} | Belum terdaftar: {n_unregistered}")
    cell.font = HEADER_FONT
    cell.fill = SUMMARY_FILL
    cell.alignment = ALIGN_LEFT
    cell.border = BORDER_THIN
    ws.row_dimensions[row].height = 18

    ws.freeze_panes = ws.cell(row=data_start, column=1)
    ws.auto_filter.ref = f"B{header_row}:M{data_end}"
    for col, width in {1: 4, 2: 5, 3: 26, 4: 36, 5: 20, 6: 12, 7: 12, 8: 12,
                       9: 14, 10: 15, 11: 12, 12: 15, 13: 42}.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    _page_setup(ws)

    # ==================== SHEET 2: REKAP PER PRODUK ====================
    ws2 = wb.create_sheet("Rekap per Produk")
    last_col2 = 9  # B..I
    _title_block(
        ws2, last_col2,
        "REKAP KELENGKAPAN DOKUMEN BAHAN BAKU PER PRODUK",
        f"Produk DR. WIDYCLINIC - {COMPANY} - {now_str}",
    )
    row = 4
    row = _table_header(ws2, row, [
        "No", "Nama Produk", "Jumlah Bahan", "Lengkap (5/5)", "Belum Lengkap",
        "Belum Terdaftar", "% Lengkap", "Bahan yang Dokumennya Belum Lengkap",
    ])
    for num in sorted(products):
        prod_name = products[num]
        used_mats = [r for r in results if num in r["used_in"]]
        complete = sum(1 for r in used_mats if r["complete"])
        partial = sum(1 for r in used_mats if r["rm"] and not r["complete"])
        unreg = sum(1 for r in used_mats if not r["rm"])
        pct = round(complete / len(used_mats) * 100) if used_mats else 0
        missing_list = []
        for r in used_mats:
            if r["complete"]:
                continue
            if r["rm"]:
                missing_list.append(f"{r['nama']} ({', '.join(r['missing'])})")
            else:
                missing_list.append(f"{r['nama']} (belum terdaftar)")

        _body_cell(ws2, row, 2, num, align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws2, row, 3, prod_name, bold=True)
        _body_cell(ws2, row, 4, len(used_mats), align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws2, row, 5, complete, align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws2, row, 6, partial, align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws2, row, 7, unreg, align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws2, row, 8, f"{pct}%", align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws2, row, 9, "; ".join(missing_list) if missing_list else "-",
                   align=ALIGN_LEFT_TOP)
        ws2.row_dimensions[row].height = 16
        row += 1

    for col, width in {1: 4, 2: 5, 3: 44, 4: 14, 5: 14, 6: 14, 7: 14, 8: 11, 9: 90}.items():
        ws2.column_dimensions[get_column_letter(col)].width = width
    ws2.freeze_panes = ws2.cell(row=5, column=1)
    _page_setup(ws2)

    # ==================== SHEET 3: MATRIKS BAHAN x PRODUK ====================
    ws3 = wb.create_sheet("Matriks Bahan-Produk")
    max_num = max(products) if products else 0
    # A=No, B=Nama Bahan, C..(2+max_num)=produk, lalu Status + Keterangan.
    mat_cols = 4 + max_num
    _title_block(
        ws3, mat_cols,
        "MATRIKS BAHAN BAKU x PRODUK + STATUS KELENGKAPAN DOKUMEN",
        f"Produk DR. WIDYCLINIC - {COMPANY} - {now_str}",
    )
    header = ["No", "Nama Bahan"] + [str(n) for n in range(1, max_num + 1)] + [
        "Status Dokumen", "Keterangan (Daftar Produk)",
    ]
    row = 4
    row = _table_header(ws3, row, header)
    for r in results:
        _body_cell(ws3, row, 1, r["no"], align=ALIGN_CENTER_NOWRAP)
        _body_cell(ws3, row, 2, r["nama"], bold=True)
        for num in range(1, max_num + 1):
            used = num in r["used_in"]
            cell = ws3.cell(row=row, column=2 + num, value="X" if used else None)
            cell.font = BODY_BOLD_FONT if used else BODY_FONT
            cell.alignment = ALIGN_CENTER_NOWRAP
            cell.border = BORDER_THIN
            if used:
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
        if not r["rm"]:
            cell = _body_cell(ws3, row, 3 + max_num, "Belum Terdaftar", align=ALIGN_CENTER_NOWRAP)
            cell.fill = NA_FILL
            cell.font = NA_FONT
        elif r["complete"]:
            cell = _body_cell(ws3, row, 3 + max_num, "Lengkap", align=ALIGN_CENTER_NOWRAP)
            cell.fill = OK_FILL
            cell.font = OK_FONT
        else:
            cell = _body_cell(ws3, row, 3 + max_num, "Belum Lengkap", align=ALIGN_CENTER_NOWRAP)
            cell.fill = BAD_FILL
            cell.font = BAD_FONT
        ket = "; ".join(f"{n}. {products[n]}" for n in r["used_in"] if n in products)
        _body_cell(ws3, row, 4 + max_num, ket or "-", align=ALIGN_LEFT)
        ws3.row_dimensions[row].height = 16
        row += 1

    ws3.column_dimensions["A"].width = 5
    ws3.column_dimensions["B"].width = 26
    for num in range(1, max_num + 1):
        ws3.column_dimensions[get_column_letter(2 + num)].width = 4.5
    ws3.column_dimensions[get_column_letter(3 + max_num)].width = 16
    ws3.column_dimensions[get_column_letter(4 + max_num)].width = 70
    ws3.freeze_panes = ws3.cell(row=5, column=1)
    _page_setup(ws3)

    wb.save(output_path)

    # ==================== RINGKASAN KONSOL ====================
    print("=" * 70)
    print("HASIL CEK KELENGKAPAN DOKUMEN BAHAN BAKU - DR. WIDYCLINIC")
    print("=" * 70)
    print(f"Perusahaan        : {COMPANY}")
    print(f"Jumlah produk     : {len(products)}")
    if n_extra_prod:
        print(f"Produk tambahan   : {n_extra_prod} dari formula sistem "
              f"({n_extra_mat} bahan baru ditambahkan)")
    print(f"Jumlah bahan baku : {total}")
    print(f"Dokumen lengkap  : {n_complete}")
    print(f"Dokumen belum    : {n_partial}")
    print(f"Belum terdaftar  : {n_unregistered}")
    print("-" * 70)
    for key, label in DOC_LABELS:
        ok_count = sum(1 for r in results if r["rm"] and r["status"][key])
        print(f"  {label:<20}: {ok_count}/{total} ada, {per_doc_missing[label]} belum")
    print("-" * 70)
    print(f"File disimpan     : {output_path}")


def main():
    source = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SOURCE
    output = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT
    if not os.path.isfile(source):
        raise SystemExit(f"File sumber tidak ditemukan: {source}")
    build_report(source, output)


if __name__ == "__main__":
    main()
